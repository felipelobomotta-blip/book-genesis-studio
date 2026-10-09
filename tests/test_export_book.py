"""Focused tests for the dependency-free Book Genesis delivery exporter."""

from pathlib import Path
import hashlib
import json
import shutil
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch


REPO_ROOT = Path(__file__).resolve().parents[1]
SCRIPT_ROOT = REPO_ROOT / "skills" / "book-genesis" / "scripts"
sys.path.insert(0, str(SCRIPT_ROOT))

from export_book import ExportError, export_book  # noqa: E402
import export_book as exporter  # noqa: E402

try:
    import docx  # noqa: F401
    import pypdf  # noqa: F401
    import reportlab  # noqa: F401
    OPTIONAL_FORMATS_AVAILABLE = True
except ImportError:
    OPTIONAL_FORMATS_AVAILABLE = False

if OPTIONAL_FORMATS_AVAILABLE:
    try:
        exporter._find_serif_font()
    except ExportError:
        OPTIONAL_FORMATS_AVAILABLE = False


class ExportBookTests(unittest.TestCase):
    def setUp(self) -> None:
        self.root = Path(tempfile.mkdtemp(prefix="book-genesis-export-"))
        self.chapters = self.root / "manuscript" / "chapters"
        self.chapters.mkdir(parents=True)
        self.delivery = self.root / "delivery"

    def tearDown(self) -> None:
        shutil.rmtree(self.root, ignore_errors=True)

    def write_chapter(self, filename: str, title: str, body: str) -> None:
        (self.chapters / filename).write_text(f"# {title}\n\n{body}\n", encoding="utf-8")

    def export_two_chapters(self) -> dict[str, object]:
        self.write_chapter("chapter-01.md", "First", "The *old* clock kept **time**.\n\n* * *\n\nThe room waited.")
        self.write_chapter("chapter-02.md", "Second", "A new morning arrived.")
        return export_book(
            self.root,
            self.delivery,
            title="Test Book",
            author="Test Author",
            language="en",
        )

    def write_complete_pipeline_evidence(self) -> None:
        artifact_names = [
            "00-brief",
            "01-market-map",
            "02-story-engine",
            "03-characters",
            "04-theme",
            "05-voice",
            "06-emotional-curve",
            "07-outline",
            "08-opening-strategy",
            "09-continuity-ledger",
            "10-adversarial-audit",
            "11-genesis-score",
            "12-editorial-package",
            "13-positioning",
        ]
        for name in artifact_names:
            (self.root / "artifacts" / f"{name}.md").parent.mkdir(parents=True, exist_ok=True)
            (self.root / "artifacts" / f"{name}.md").write_text(
                "audit_status: major_rewrite\n" if name == "10-adversarial-audit" else f"evidence for {name}\n",
                encoding="utf-8",
            )
        for relative in (
            "ASSUMPTIONS.md",
            "RUN_REPORT.md",
            "evaluations/panel-chapter-01.md",
            "evaluations/proofread.md",
            "evaluations/revision-loop.md",
            "evaluations/revision-plan.md",
        ):
            path = self.root / relative
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(f"evidence for {relative}\n", encoding="utf-8")
        (self.root / "PROJECT_STATE.yaml").write_text(
            """schema_version: 6

pipeline:
  current_phase: "Phase 7: Editorial Package"
  status: "in_progress"

manuscript:
  chapter_count_planned: 1
  completed_chapters: [1]
  word_count_actual: 3
  status: "complete"

gates:
  intake: "passed"
  foundation: "passed"
  architecture: "passed"
  drafting: "passed"
  adversarial_audit: "failed"
  revision_loop: "failed"
  final_score: "failed"
  editorial_package: "in_progress"
""",
            encoding="utf-8",
        )

    def test_gaps_and_duplicate_numbers_are_rejected(self) -> None:
        self.write_chapter("chapter-01.md", "First", "One.")
        self.write_chapter("chapter-03.md", "Third", "Three.")
        with self.assertRaisesRegex(ExportError, "gaps.*02"):
            exporter.discover_chapters(self.chapters)

        (self.chapters / "chapter-03.md").unlink()
        self.write_chapter("chapter-1.md", "Duplicate", "One again.")
        with self.assertRaisesRegex(ExportError, "duplicate chapter number 1"):
            exporter.discover_chapters(self.chapters)

    def test_manuscript_and_epub_preserve_order_text_italics_and_toc(self) -> None:
        result = self.export_two_chapters()
        manuscript = self.delivery / "manuscript.md"
        epub = self.delivery / "test-book.epub"
        receipt_path = self.delivery / "export-receipt.json"

        self.assertEqual(2, result["chapter_count"])
        text = manuscript.read_text(encoding="utf-8")
        self.assertLess(text.index("# First"), text.index("# Second"))
        self.assertIn("The *old* clock kept **time**.", text)
        receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
        self.assertEqual(2, receipt["chapter_count"])
        self.assertEqual("whitespace_split_prose_v1", receipt["word_count"]["method"])
        self.assertEqual(
            [{"number": 1, "words": 8}, {"number": 2, "words": 4}],
            receipt["word_count"]["chapters"],
        )
        self.assertEqual(12, receipt["word_count"]["total"])
        self.assertEqual("not_assessed", receipt["quality"]["status"])
        self.assertEqual(hashlib.sha256((self.chapters / "chapter-01.md").read_bytes()).hexdigest(), receipt["source_chapters"][0]["sha256"])
        self.assertEqual(hashlib.sha256(manuscript.read_bytes()).hexdigest(), receipt["outputs"]["manuscript.md"]["sha256"])
        self.assertEqual("passed", receipt["validation"]["status"])

        with zipfile.ZipFile(epub) as archive:
            self.assertEqual("mimetype", archive.namelist()[0])
            self.assertEqual(zipfile.ZIP_STORED, archive.getinfo("mimetype").compress_type)
            chapter_one = archive.read("EPUB/text/chapter-01.xhtml").decode("utf-8")
            nav = archive.read("EPUB/nav.xhtml").decode("utf-8")
            self.assertIn("<em>old</em>", chapter_one)
            self.assertIn('<hr class="scene-break" />', chapter_one)
            self.assertLess(nav.index("chapter-01.xhtml"), nav.index("chapter-02.xhtml"))
            opf = archive.read("EPUB/content.opf").decode("utf-8")
            self.assertIn("title-page", opf)
            self.assertIn("chapter-01", opf)
            self.assertIn("<p>Test Author</p>", archive.read("EPUB/titlepage.xhtml").decode("utf-8"))

    def test_nonempty_delivery_requires_deliberate_overwrite_and_failure_preserves_it(self) -> None:
        self.export_two_chapters()
        old_manuscript = (self.delivery / "manuscript.md").read_bytes()
        old_receipt = (self.delivery / "export-receipt.json").read_bytes()
        with self.assertRaisesRegex(ExportError, "not empty"):
            export_book(self.root, self.delivery, title="Test Book", author="Test Author", language="en")
        self.assertEqual(old_manuscript, (self.delivery / "manuscript.md").read_bytes())

        with patch.object(exporter, "write_epub", side_effect=OSError("injected export failure")):
            with self.assertRaisesRegex(OSError, "injected export failure"):
                export_book(self.root, self.delivery, title="Test Book", author="Test Author", language="en", overwrite=True)
        self.assertEqual(old_manuscript, (self.delivery / "manuscript.md").read_bytes())
        self.assertEqual(old_receipt, (self.delivery / "export-receipt.json").read_bytes())
        self.assertEqual([], list(self.root.glob(".delivery.stage-*")))
        self.assertEqual([], list(self.root.glob(".delivery.backup-*")))

    def test_slug_and_delivery_target_cannot_escape_book_root(self) -> None:
        self.write_chapter("chapter-01.md", "First", "One.")
        with self.assertRaisesRegex(ExportError, "simple lowercase filename"):
            export_book(self.root, self.delivery, title="Test Book", author="Test Author", language="en", slug="../outside")
        with self.assertRaisesRegex(ExportError, "book_dir/delivery"):
            export_book(self.root, self.root / "unsafe-delivery", title="Test Book", author="Test Author", language="en")
        with self.assertRaisesRegex(ExportError, "book_dir/delivery"):
            export_book(self.root, self.root, title="Test Book", author="Test Author", language="en")

    def test_complete_pipeline_missing_audit_fails_before_delivery(self) -> None:
        self.write_chapter("chapter-01.md", "First", "One two three.")
        self.write_complete_pipeline_evidence()
        (self.root / "artifacts" / "10-adversarial-audit.md").unlink()

        with self.assertRaisesRegex(ExportError, "10-adversarial-audit.md"):
            export_book(
                self.root,
                self.delivery,
                title="Test Book",
                author="Test Author",
                language="en",
                require_complete_pipeline=True,
                expected_chapters=1,
                min_words=1,
                max_words=10,
            )
        self.assertFalse(self.delivery.exists())

    def test_complete_pipeline_missing_proofread_fails_before_delivery(self) -> None:
        self.write_chapter("chapter-01.md", "First", "One two three.")
        self.write_complete_pipeline_evidence()
        (self.root / "evaluations" / "proofread.md").unlink()

        with self.assertRaisesRegex(ExportError, "evaluations/proofread.md"):
            export_book(
                self.root,
                self.delivery,
                title="Test Book",
                author="Test Author",
                language="en",
                require_complete_pipeline=True,
                expected_chapters=1,
                min_words=1,
                max_words=10,
            )
        self.assertFalse(self.delivery.exists())

    def test_complete_pipeline_allows_honest_failed_audit_report(self) -> None:
        self.write_chapter("chapter-01.md", "First", "One two three.")
        self.write_complete_pipeline_evidence()

        result = export_book(
            self.root,
            self.delivery,
            title="Test Book",
            author="Test Author",
            language="en",
            require_complete_pipeline=True,
            expected_chapters=1,
            min_words=3,
            max_words=3,
        )

        self.assertTrue(Path(result["receipt"]).is_file())

    def test_complete_pipeline_rejects_skipped_audit_even_with_nonempty_evidence(self) -> None:
        self.write_chapter("chapter-01.md", "First", "One two three.")
        self.write_complete_pipeline_evidence()
        state_path = self.root / "PROJECT_STATE.yaml"
        state_path.write_text(
            state_path.read_text(encoding="utf-8").replace('adversarial_audit: "failed"', 'adversarial_audit: "skipped"'),
            encoding="utf-8",
        )

        with self.assertRaisesRegex(ExportError, "gates.adversarial_audit"):
            export_book(
                self.root,
                self.delivery,
                title="Test Book",
                author="Test Author",
                language="en",
                require_complete_pipeline=True,
                expected_chapters=1,
                min_words=1,
                max_words=10,
            )
        self.assertFalse(self.delivery.exists())

    def test_complete_pipeline_rejects_chapter_and_word_contract_mismatch(self) -> None:
        self.write_chapter("chapter-01.md", "First", "One two three.")
        self.write_complete_pipeline_evidence()
        kwargs = {
            "require_complete_pipeline": True,
            "min_words": 1,
            "max_words": 10,
        }
        with self.assertRaisesRegex(ExportError, "expected 2 chapters but found 1"):
            export_book(
                self.root,
                self.delivery,
                title="Test Book",
                author="Test Author",
                language="en",
                expected_chapters=2,
                **kwargs,
            )
        with self.assertRaisesRegex(ExportError, "below min_words 4"):
            export_book(
                self.root,
                self.delivery,
                title="Test Book",
                author="Test Author",
                language="en",
                expected_chapters=1,
                min_words=4,
                max_words=10,
                require_complete_pipeline=True,
            )
        self.assertFalse(self.delivery.exists())

    def test_source_change_before_promotion_preserves_previous_delivery(self) -> None:
        self.export_two_chapters()
        old_files = {
            path.name: path.read_bytes()
            for path in self.delivery.iterdir()
            if path.is_file()
        }
        original_validate = exporter.validate_epub

        def mutate_after_validation(path: Path, expected_chapters: int) -> dict[str, object]:
            result = original_validate(path, expected_chapters)
            (self.chapters / "chapter-02.md").write_text("# Second\n\nChanged during export.\n", encoding="utf-8")
            return result

        with patch.object(exporter, "validate_epub", side_effect=mutate_after_validation):
            with self.assertRaisesRegex(ExportError, "source chapter set changed"):
                export_book(self.root, self.delivery, title="Test Book", author="Test Author", language="en", overwrite=True)
        self.assertEqual(old_files, {path.name: path.read_bytes() for path in self.delivery.iterdir() if path.is_file()})
        self.assertEqual([], list(self.root.glob(".delivery.stage-*")))
        self.assertEqual([], list(self.root.glob(".delivery.backup-*")))

    def test_source_change_after_pre_promotion_check_rolls_back_promoted_delivery(self) -> None:
        self.export_two_chapters()
        old_files = {
            path.name: path.read_bytes()
            for path in self.delivery.iterdir()
            if path.is_file()
        }
        original_assert = exporter._assert_sources_unchanged
        checks = 0

        def mutate_after_first_check(chapters_dir: Path, chapters: list[object]) -> None:
            nonlocal checks
            original_assert(chapters_dir, chapters)
            checks += 1
            if checks == 1:
                (self.chapters / "chapter-02.md").write_text("# Second\n\nChanged after staging.\n", encoding="utf-8")

        with patch.object(exporter, "_assert_sources_unchanged", side_effect=mutate_after_first_check):
            with self.assertRaisesRegex(ExportError, "source chapter set changed"):
                export_book(self.root, self.delivery, title="Test Book", author="Test Author", language="en", overwrite=True)
        self.assertEqual(old_files, {path.name: path.read_bytes() for path in self.delivery.iterdir() if path.is_file()})
        self.assertFalse((self.chapters / ".book-genesis-write.lock").exists())
        self.assertEqual([], list(self.root.glob(".delivery.stage-*")))
        self.assertEqual([], list(self.root.glob(".delivery.backup-*")))

    def test_book_lock_blocks_a_cooperating_concurrent_writer(self) -> None:
        self.write_chapter("chapter-01.md", "First", "One.")
        lock = self.chapters / ".book-genesis-write.lock"
        lock.write_text("pid=other\n", encoding="ascii")
        try:
            with self.assertRaisesRegex(ExportError, "book is locked"):
                export_book(self.root, self.delivery, title="Test Book", author="Test Author", language="en")
        finally:
            lock.unlink()

    def test_xml_invalid_control_characters_are_rejected_in_source_and_epub(self) -> None:
        self.write_chapter("chapter-01.md", "One", "hello\x00world")
        with self.assertRaisesRegex(ExportError, "XML-invalid"):
            exporter.discover_chapters(self.chapters)

        chapter = exporter.Chapter(1, self.chapters / "chapter-01.md", "One", "hello\x00world\n", "0" * 64)
        metadata = exporter.BookMetadata("Test Book", "Test Author", "en-US", "test-book")
        bad_epub = self.root / "bad.epub"
        exporter.write_epub(bad_epub, [chapter], metadata)
        with self.assertRaisesRegex(ExportError, "XHTML is not well-formed"):
            exporter.validate_epub(bad_epub, 1)

    @unittest.skipUnless(OPTIONAL_FORMATS_AVAILABLE, "optional document/PDF dependencies are unavailable")
    def test_blank_pdf_is_rejected_by_full_text_validation(self) -> None:
        from reportlab.pdfgen.canvas import Canvas

        blank_pdf = self.root / "blank.pdf"
        canvas = Canvas(str(blank_pdf), pagesize=(396, 612))
        for _ in range(3):
            canvas.showPage()
        canvas.save()
        with self.assertRaisesRegex(ExportError, "no readable prose"):
            exporter.validate_pdf(blank_pdf, 2)

    @unittest.skipUnless(OPTIONAL_FORMATS_AVAILABLE, "optional document/PDF dependencies are unavailable")
    def test_optional_docx_and_pdf_preserve_text_and_are_validated(self) -> None:
        self.write_chapter("chapter-01.md", "First", "The *old* clock kept **time**.\n\n* * *\n\nThe room waited.")
        self.write_chapter("chapter-02.md", "Second", "A new morning arrived.")
        result = export_book(self.root, self.delivery, title="Test Book", author="Test Author", language="en", docx=True, pdf=True)
        self.assertTrue(Path(result["docx"]).is_file())
        self.assertTrue(Path(result["pdf"]).is_file())
        receipt = json.loads((self.delivery / "export-receipt.json").read_text(encoding="utf-8"))
        self.assertEqual("passed", receipt["validation"]["docx"]["status"])
        self.assertEqual("passed", receipt["validation"]["pdf"]["status"])

        from docx import Document
        document = Document(self.delivery / "test-book.docx")
        docx_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("The old clock kept time.", docx_text)
        self.assertTrue(any(run.italic for paragraph in document.paragraphs for run in paragraph.runs if "old" in run.text))
        self.assertTrue(any(run.bold for paragraph in document.paragraphs for run in paragraph.runs if "time" in run.text))

        from pypdf import PdfReader
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(self.delivery / "test-book.pdf").pages)
        self.assertIn("The old clock kept time.", pdf_text)
        self.assertIn("A new morning arrived.", pdf_text)


if __name__ == "__main__":
    unittest.main()
